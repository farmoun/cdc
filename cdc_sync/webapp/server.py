"""CDC 监控 Web 应用（FastAPI）。

只读端点 + 一键操作端点，全部复用 cdc_sync 既有模块。
任一后端不可达时返回结构化 error，不 500 崩溃。

配置来源：环境变量 CDC_SETTINGS / CDC_TABLES 指定 settings/tables 路径；
默认 config/settings.yaml 与 config/tables.yaml。
"""
from __future__ import annotations

import logging
import os
from pathlib import Path

from fastapi import FastAPI, Request
from fastapi.responses import FileResponse, JSONResponse, RedirectResponse, Response

from . import auth

from .. import (
    ck_client,
    ck_generator,
    config,
    config_store,
    connect_client,
    connector_generator,
    query_store,
    reconcile as reconcile_mod,
    sql_parser,
    ui_state,
)
from . import conn_test, kafka_lag

log = logging.getLogger("cdc_sync.webapp")

STATIC_DIR = Path(__file__).resolve().parent / "static"

app = FastAPI(title="CDC 同步监控", version="1.0.0")


@app.middleware("http")
async def _no_cache_static(request: Request, call_next):
    """静态资源(html/js/css)禁用缓存，避免改了前端浏览器还用旧版。"""
    resp = await call_next(request)
    path = request.url.path
    if path == "/" or path.endswith((".js", ".css", ".html")):
        resp.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
        resp.headers["Pragma"] = "no-cache"
    return resp


@app.middleware("http")
async def _require_login(request: Request, call_next):
    """除公开端点外的所有 /api/* 需携带有效会话 Cookie，否则返回 401。

    /api/me 也放行到 handler：handler 内部按有无有效会话返回 me 或 401，
    前端据此决定显示登录框还是主界面。
    """
    path = request.url.path
    if not path.startswith("/api/"):
        return await call_next(request)  # 静态页面等无需鉴权
    if path in auth.PUBLIC_API_PATHS:
        return await call_next(request)
    token = request.cookies.get(auth.COOKIE_NAME)
    if not token or not auth.validate_token(token):
        return JSONResponse({"ok": False, "error": "未登录或会话已过期", "auth": True}, status_code=401)
    return await call_next(request)


def _settings_path() -> str | None:
    return os.environ.get("CDC_SETTINGS") or None


def _tables_path() -> str | None:
    return os.environ.get("CDC_TABLES") or None


def _load():
    settings = config.load_settings(_settings_path())
    tables = config.load_tables(_tables_path())
    return settings, tables


def _ok(data):
    return {"ok": True, "data": data}


def _err(msg: str):
    return {"ok": False, "error": str(msg)}


# ----------------------------- 登录鉴权 -----------------------------

@app.get("/api/health")
def api_health():
    """无鉴权健康检查端点(供 Docker HEALTHCHECK 用)，恒返回 ok。"""
    return _ok({"status": "ok"})


@app.get("/api/me")
def api_me(request: Request):
    """返回当前登录用户；未登录返回 401，前端据此切换登录/主界面。"""
    token = request.cookies.get(auth.COOKIE_NAME)
    username = auth.validate_token(token) if token else None
    if not username:
        return JSONResponse({"ok": False, "error": "未登录", "auth": True}, status_code=401)
    return _ok({"username": username, "session_seconds": auth.SESSION_SECONDS})


@app.post("/api/login")
async def api_login(request: Request):
    """校验用户名+密码，成功则签发会话 Cookie。body: {username, password}"""
    try:
        body = await request.json()
    except Exception as e:  # noqa: BLE001
        return JSONResponse({"ok": False, "error": f"请求体非法 JSON：{e}", "auth": True}, status_code=400)
    username = (body.get("username") or "").strip()
    password = body.get("password") or ""
    if not auth.verify_password(username, password):
        # 取真实客户端 IP（支持反代 X-Forwarded-For）
        forwarded = request.headers.get("X-Forwarded-For")
        client_ip = forwarded.split(",")[0].strip() if forwarded else (
            request.client.host if request.client else "unknown"
        )
        auth.record_fail(client_ip)
        return JSONResponse({"ok": False, "error": "用户名或密码错误", "auth": True}, status_code=401)
    token = auth.issue_token(username)
    resp = JSONResponse({"ok": True, "data": {"username": username}})
    resp.set_cookie(
        auth.COOKIE_NAME, token,
        max_age=auth.SESSION_SECONDS, httponly=True, samesite="lax",
        secure=auth.COOKIE_SECURE, path="/",
    )
    return resp


@app.post("/api/logout")
def api_logout():
    """清除会话 Cookie。"""
    resp = JSONResponse({"ok": True, "data": {"logged_out": True}})
    resp.delete_cookie(auth.COOKIE_NAME, path="/")
    return resp


@app.get("/api/alerts")
def api_alerts():
    """返回告警列表（最新在前）。需要已登录。"""
    return _ok(auth.get_alerts())


@app.post("/api/alerts/record")
async def api_record_alert(request: Request):
    """前端主动记录一条操作告警（如用户确认执行高风险 SQL 后调用）。"""
    import time
    try:
        body = await request.json()
    except Exception as e:  # noqa: BLE001
        return _err(f"请求体非法 JSON：{e}")
    ip = (request.client.host if request.client else "") or ""
    token = request.cookies.get(auth.COOKIE_NAME, "")
    username = auth.validate_token(token) if token else ""
    entry = {
        "time": time.strftime("%Y-%m-%d %H:%M:%S", time.localtime()),
        "type": body.get("type", "risky_query"),
        "action": body.get("action", "confirmed_execute"),
        "ip": ip,
        "user": username or "—",
        "detail": body.get("detail", ""),
        "sql": (body.get("sql") or "")[:500],
    }
    auth.record_alert(entry)
    return _ok({"recorded": True})


@app.post("/api/ai/check")
async def api_ai_check(request: Request):
    """用 AI 检测 SQL 风险。配置 settings.yaml 的 ai 段后生效；未配置时跳过(返回 skipped)。"""
    import json as _json
    import urllib.request as _req
    import urllib.error
    try:
        body = await request.json()
    except Exception as e:  # noqa: BLE001
        return _err(f"请求体非法 JSON：{e}")
    sql = (body.get("sql") or "").strip()
    if not sql:
        return _err("sql 为空")
    try:
        s = config_store.read_settings_dict()
        ai_cfg = s.get("ai") or {}
        base_url = (ai_cfg.get("url") or "").rstrip("/")
        api_key = ai_cfg.get("key") or ""
        model = ai_cfg.get("model") or "gpt-4o-mini"
    except Exception as e:  # noqa: BLE001
        return _err(f"配置读取失败：{e}")
    if not base_url or not api_key:
        return _ok({"risk": False, "skipped": True, "reason": "AI 未配置"})
    prompt = (
        "你是一个数据库安全审计助手。请严格按照以下规则判断 SQL 语句是否属于高风险操作。\n\n"
        "【高风险（risk=true）】——以下任何一条满足即为高风险，必须返回 risk=true：\n"
        "1. 增：INSERT、REPLACE、LOAD DATA、IMPORT\n"
        "2. 删：DELETE、DROP、TRUNCATE\n"
        "3. 改：UPDATE、ALTER、RENAME、MODIFY\n"
        "4. 权限与结构：CREATE、GRANT、REVOKE、CALL、EXECUTE、MERGE\n"
        "5. 任何会导致数据写入、数据删除、表结构变更、权限变更的操作\n"
        "6. 批量操作或不带 WHERE 条件的 UPDATE/DELETE\n\n"
        "【低风险（risk=false）】——仅以下情况才是低风险：\n"
        "纯只读查询：SELECT（不含子查询写操作）、SHOW、EXPLAIN、DESCRIBE\n\n"
        "判断原则：宁可误报，不可漏报。只要语句涉及增、删、改或其他任何非只读操作，一律返回 risk=true。\n\n"
        "请只返回如下 JSON，不要有任何多余内容：\n"
        '{"risk": true/false, "level": "high"/"low", "reason": "一句话说明风险原因"}\n\n'
        f"SQL：\n{sql}"
    )
    payload = _json.dumps({
        "model": model,
        "messages": [{"role": "user", "content": prompt}],
        "max_tokens": 120,
        "temperature": 0,
    }).encode("utf-8")
    api_url = f"{base_url}/chat/completions"
    http_req = _req.Request(
        api_url,
        data=payload,
        headers={
            "Content-Type": "application/json",
            "Authorization": f"Bearer {api_key}",
        },
        method="POST",
    )
    try:
        with _req.urlopen(http_req, timeout=15) as resp:
            resp_body = _json.loads(resp.read().decode("utf-8"))
        text = resp_body["choices"][0]["message"]["content"].strip()
        # 提取 JSON（模型可能在前后多输出内容）
        start = text.find("{")
        end = text.rfind("}") + 1
        if start >= 0 and end > start:
            result = _json.loads(text[start:end])
        else:
            result = {"risk": False, "level": "low", "reason": text}
        return _ok(result)
    except urllib.error.HTTPError as e:
        err_body = e.read().decode("utf-8", errors="replace")[:300]
        return _err(f"AI API 返回 {e.code}：{err_body}")
    except Exception as e:  # noqa: BLE001
        return _err(f"AI 请求失败：{e}")

@app.get("/api/connectors")
def api_connectors():
    try:
        settings, _ = _load()
    except Exception as e:  # noqa: BLE001
        return _err(e)
    name = settings.debezium.connector_name
    try:
        st = connect_client.status(settings.connect_url, name)
        return _ok(st)
    except Exception as e:  # noqa: BLE001
        return _err(e)


@app.get("/api/kafka/lag")
def api_kafka_lag():
    try:
        settings, tables = _load()
    except Exception as e:  # noqa: BLE001
        return _err(e)
    try:
        rows = kafka_lag.topic_lag(settings, tables.tables)
        return _ok(rows)
    except Exception as e:  # noqa: BLE001
        return _err(e)


@app.get("/api/clickhouse/consumers")
def api_ck_consumers():
    try:
        settings, _ = _load()
    except Exception as e:  # noqa: BLE001
        return _err(e)
    try:
        rows = ck_client.kafka_consumers_status(settings.clickhouse)
        return _ok(rows)
    except Exception as e:  # noqa: BLE001
        return _err(e)


@app.get("/api/reconcile")
def api_reconcile():
    try:
        settings, tables = _load()
    except Exception as e:  # noqa: BLE001
        return _err(e)
    try:
        results = reconcile_mod.reconcile_all(settings.mysql, settings.clickhouse, tables.tables)
        data = [
            {"table": r.table, "mysql": r.mysql_count, "clickhouse": r.ck_count,
             "diff": r.diff, "consistent": r.consistent}
            for r in results
        ]
        return _ok(data)
    except Exception as e:  # noqa: BLE001
        return _err(e)


@app.get("/api/overview")
def api_overview():
    """聚合各组件健康：connect / kafka / clickhouse / config。"""
    result = {"components": {}}
    try:
        settings, tables = _load()
        result["config"] = {"status": "ok", "tables": len(tables.tables),
                            "connector": settings.debezium.connector_name}
    except Exception as e:  # noqa: BLE001
        result["config"] = {"status": "error", "error": str(e)}
        return result

    # Connect
    try:
        st = connect_client.status(settings.connect_url, settings.debezium.connector_name)
        cstate = st.get("connector", {}).get("state", "?")
        tasks = st.get("tasks", [])
        running = sum(1 for t in tasks if t.get("state") == "RUNNING")
        result["components"]["connect"] = {
            "status": "ok" if cstate == "RUNNING" else "warn",
            "state": cstate, "tasks_total": len(tasks), "tasks_running": running,
        }
    except Exception as e:  # noqa: BLE001
        result["components"]["connect"] = {"status": "error", "error": str(e)}

    # Kafka（总堆积）
    try:
        rows = kafka_lag.topic_lag(settings, tables.tables)
        if rows and rows[0].get("error") and len(rows) == 1:
            result["components"]["kafka"] = {"status": "error", "error": rows[0]["error"]}
        else:
            total_lag = sum(r.get("lag", 0) for r in rows if isinstance(r.get("lag"), int))
            result["components"]["kafka"] = {"status": "ok", "topics": len(rows), "total_lag": total_lag}
    except Exception as e:  # noqa: BLE001
        result["components"]["kafka"] = {"status": "error", "error": str(e)}

    # ClickHouse
    try:
        rows = ck_client.kafka_consumers_status(settings.clickhouse)
        exc = sum(1 for r in rows if r.get("last_exception"))
        result["components"]["clickhouse"] = {
            "status": "ok" if exc == 0 else "warn",
            "consumers": len(rows), "with_exception": exc,
        }
    except Exception as e:  # noqa: BLE001
        result["components"]["clickhouse"] = {"status": "error", "error": str(e)}

    return result


# ----------------------------- 操作端点 -----------------------------

@app.post("/api/connectors/deploy")
def api_deploy():
    try:
        settings, tables = _load()
    except Exception as e:  # noqa: BLE001
        return _err(e)
    try:
        connector = connector_generator.build_connector(tables.tables, settings)
        resp = connect_client.deploy(settings.connect_url, connector)
        return _ok({"name": resp.get("name", connector["name"])})
    except Exception as e:  # noqa: BLE001
        return _err(e)


@app.post("/api/connectors/restart")
def api_restart():
    try:
        settings, _ = _load()
    except Exception as e:  # noqa: BLE001
        return _err(e)
    try:
        connect_client.restart(settings.connect_url, settings.debezium.connector_name)
        return _ok({"restarted": settings.debezium.connector_name})
    except Exception as e:  # noqa: BLE001
        return _err(e)


@app.post("/api/clickhouse/apply")
def api_apply_ck():
    """对所有表执行 CK 三对象 DDL（IF NOT EXISTS，可重复）。"""
    try:
        settings, tables = _load()
    except Exception as e:  # noqa: BLE001
        return _err(e)
    applied = []
    try:
        try:
            _, ck_major = ck_client.server_version(settings.clickhouse)
        except Exception:  # noqa: BLE001
            ck_major = None
        for t in tables.tables:
            if not t.has_columns():
                return _err(f"表 {t.source_table} 无列定义，请先 import-sql/introspect")
            tsql = ck_generator.build_table_sql(t, settings, ck_major=ck_major)
            ck_client.execute_statements(settings.clickhouse, tsql.ordered(), dry_run=False)
            applied.append(t.target_table)
        return _ok({"applied": applied, "ck_major": ck_major})
    except Exception as e:  # noqa: BLE001
        return _err(f"已建 {len(applied)} 张，失败于下一张：{e}")


@app.post("/api/reconcile/run")
def api_reconcile_run():
    return api_reconcile()


# ----------------------------- 同步管道主开关 -----------------------------

@app.post("/api/pipeline/start")
def api_pipeline_start():
    """开启同步：挂 CK 消费 + 恢复/发布连接器 + (schema_only 触发增量快照)。"""
    from .. import pipeline
    try:
        settings, tables = _load()
    except Exception as e:  # noqa: BLE001
        return _err(e)
    try:
        for t in tables.tables:
            if not t.has_columns():
                return _err(f"表 {t.source_table} 无列定义，请先在配置页导入表结构")
        r = pipeline.start(settings, tables)
        return _ok(r)
    except Exception as e:  # noqa: BLE001
        return _err(e)


@app.post("/api/pipeline/stop")
def api_pipeline_stop():
    """停止同步：暂停连接器 + 摘除 CK 消费。"""
    from .. import pipeline
    try:
        settings, tables = _load()
    except Exception as e:  # noqa: BLE001
        return _err(e)
    try:
        return _ok(pipeline.stop(settings, tables))
    except Exception as e:  # noqa: BLE001
        return _err(e)


@app.post("/api/pipeline/deploy")
def api_pipeline_deploy():
    """完整部署但停止态：建表 + 摘消费 + 注册连接器 + 暂停。"""
    from .. import pipeline
    try:
        settings, tables = _load()
    except Exception as e:  # noqa: BLE001
        return _err(e)
    try:
        for t in tables.tables:
            if not t.has_columns():
                return _err(f"表 {t.source_table} 无列定义，请先在配置页导入表结构")
        try:
            _, ck_major = ck_client.server_version(settings.clickhouse)
        except Exception:  # noqa: BLE001
            ck_major = None
        pipeline.deploy_idle(settings, tables, ck_major=ck_major)
        return _ok({"deployed": len(tables.tables), "state": "stopped"})
    except Exception as e:  # noqa: BLE001
        return _err(e)


@app.get("/api/pipeline")
def api_pipeline_state():
    """返回管道期望状态(ui_state) + 连接器实时状态。"""
    st = {"desired": ui_state.load_state().get("pipeline", "stopped")}
    try:
        settings, _ = _load()
        st["connector"] = connect_client.connector_state(settings.connect_url, settings.debezium.connector_name)
    except Exception as e:  # noqa: BLE001
        st["connector"] = None
        st["error"] = str(e)
    return _ok(st)


# ----------------------------- 配置读写端点 -----------------------------

@app.get("/api/config/settings")
def api_get_settings():
    try:
        return _ok(config_store.read_settings_dict())
    except Exception as e:  # noqa: BLE001
        return _err(e)


@app.post("/api/config/settings")
async def api_save_settings(request: Request):
    try:
        data = await request.json()
    except Exception as e:  # noqa: BLE001
        return _err(f"请求体非法 JSON：{e}")
    try:
        config_store.save_settings_dict(data)
        return _ok({"saved": True, "path": str(config_store.active_settings_path())})
    except Exception as e:  # noqa: BLE001
        return _err(e)


@app.get("/api/config/tables")
def api_get_tables():
    try:
        text, count = config_store.read_tables_text()
        return _ok({"text": text, "count": count})
    except Exception as e:  # noqa: BLE001
        return _err(e)


@app.post("/api/config/tables")
async def api_save_tables(request: Request):
    try:
        body = await request.json()
        text = body.get("text", "")
    except Exception as e:  # noqa: BLE001
        return _err(f"请求体非法 JSON：{e}")
    try:
        count = config_store.save_tables_text(text)
        return _ok({"saved": True, "count": count})
    except Exception as e:  # noqa: BLE001
        return _err(e)


@app.post("/api/config/import-sql")
async def api_import_sql(request: Request):
    """上传 SQL DDL 文本，解析生成 tables.yaml。

    body: {sql: "<内容>", database: "yibuapi", target_database?: "", prefix?: ""}
    """
    try:
        body = await request.json()
    except Exception as e:  # noqa: BLE001
        return _err(f"请求体非法 JSON：{e}")
    sql = body.get("sql", "")
    database = (body.get("database") or "").strip()
    if not sql.strip():
        return _err("SQL 内容为空")
    if not database:
        return _err("请填写源库名 database")
    try:
        parsed = sql_parser.parse_text(sql)
        if not parsed:
            return _err("未解析到任何 CREATE TABLE")
        doc = sql_parser.build_tables_doc(
            parsed, database=database,
            target_database=(body.get("target_database") or "").strip() or None,
            prefix=(body.get("prefix") or "").strip(),
        )
        count = config_store.write_tables_from_parsed(doc)
        no_pk = [t["source_table"] for t in doc["tables"] if t.get("order_by") == ""]
        return _ok({"tables": count, "no_pk": no_pk})
    except Exception as e:  # noqa: BLE001
        return _err(e)


# ----------------------------- 连通性测试端点 -----------------------------

@app.post("/api/test/{component}")
async def api_test(component: str, request: Request):
    """测试某组件连通性。body 可选传 settings 表单值（未保存也能测）；否则用已保存配置。"""
    try:
        body = await request.json()
    except Exception:  # noqa: BLE001
        body = None
    try:
        if isinstance(body, dict) and body:
            settings = config.build_settings(body)
        else:
            settings = config.load_settings(_settings_path())
    except Exception as e:  # noqa: BLE001
        return _err(f"配置解析失败：{e}")
    try:
        return _ok(conn_test.run_test(component, settings))
    except Exception as e:  # noqa: BLE001
        return _err(e)


# ----------------------------- 便捷查询端点 -----------------------------

@app.post("/api/query")
async def api_query(request: Request):
    """对 ClickHouse 执行任意 SQL，返回列/行（截断 1000 行）。"""
    try:
        body = await request.json()
    except Exception as e:  # noqa: BLE001
        return _err(f"请求体非法 JSON：{e}")
    sql = (body.get("sql") or "").strip()
    if not sql:
        return _err("查询语句为空")
    try:
        settings = config.load_settings(_settings_path())
    except Exception as e:  # noqa: BLE001
        return _err(f"配置加载失败：{e}")
    try:
        return _ok(ck_client.run_query(settings.clickhouse, sql))
    except Exception as e:  # noqa: BLE001
        return _err(e)


def _export_bytes(columns: list, rows: list, fmt: str):
    """把查询结果导出为 csv/xlsx/xls 字节流。返回 (data, media_type, ext)。"""
    import io
    fmt = (fmt or "csv").lower()
    if fmt == "csv":
        import csv
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(columns)
        w.writerows(rows)
        return buf.getvalue().encode("utf-8-sig"), "text/csv; charset=utf-8", "csv"  # BOM 让 Excel 正确显示中文
    if fmt == "xlsx":
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(list(columns))
        for r in rows:
            ws.append([v for v in r])
        bio = io.BytesIO()
        wb.save(bio)
        return bio.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
    if fmt == "xls":
        import xlwt
        wb = xlwt.Workbook(encoding="utf-8")
        ws = wb.add_sheet("result")
        for c, col in enumerate(columns):
            ws.write(0, c, str(col))
        for ri, r in enumerate(rows, start=1):
            for c, v in enumerate(r):
                ws.write(ri, c, v if isinstance(v, (int, float, str)) else ("" if v is None else str(v)))
        bio = io.BytesIO()
        wb.save(bio)
        return bio.getvalue(), "application/vnd.ms-excel", "xls"
    raise ValueError(f"不支持的导出格式：{fmt}（支持 csv/xlsx/xls）")


@app.post("/api/query/export")
async def api_query_export(request: Request):
    """执行 SQL 并导出结果为 csv/xlsx/xls 文件。"""
    try:
        body = await request.json()
    except Exception as e:  # noqa: BLE001
        return JSONResponse({"ok": False, "error": f"请求体非法 JSON：{e}"}, status_code=400)
    sql = (body.get("sql") or "").strip()
    fmt = (body.get("format") or "csv").lower()
    if not sql:
        return JSONResponse({"ok": False, "error": "查询语句为空"}, status_code=400)
    try:
        settings = config.load_settings(_settings_path())
        res = ck_client.run_query(settings.clickhouse, sql, max_rows=100000)  # 导出放宽行数上限
        data, media, ext = _export_bytes(res["columns"], res["rows"], fmt)
    except Exception as e:  # noqa: BLE001
        return JSONResponse({"ok": False, "error": str(e)}, status_code=400)
    return Response(
        content=data, media_type=media,
        headers={"Content-Disposition": f'attachment; filename="query_export.{ext}"'},
    )


@app.get("/api/queries")
def api_list_queries():
    try:
        return _ok(query_store.load_queries())
    except Exception as e:  # noqa: BLE001
        return _err(e)


@app.post("/api/queries")
async def api_save_query(request: Request):
    try:
        body = await request.json()
    except Exception as e:  # noqa: BLE001
        return _err(f"请求体非法 JSON：{e}")
    try:
        items = query_store.save_query(body.get("name", ""), body.get("sql", ""))
        return _ok(items)
    except Exception as e:  # noqa: BLE001
        return _err(e)


@app.delete("/api/queries/{name}")
def api_delete_query(name: str):
    try:
        return _ok(query_store.delete_query(name))
    except Exception as e:  # noqa: BLE001
        return _err(e)


# ----------------------------- UI 状态（监控开关等）-----------------------------

@app.get("/api/ui-state")
def api_get_ui_state():
    try:
        return _ok(ui_state.load_state())
    except Exception as e:  # noqa: BLE001
        return _err(e)


@app.post("/api/ui-state")
async def api_set_ui_state(request: Request):
    try:
        body = await request.json()
    except Exception as e:  # noqa: BLE001
        return _err(f"请求体非法 JSON：{e}")
    try:
        return _ok(ui_state.save_state(body))
    except Exception as e:  # noqa: BLE001
        return _err(e)


# ----------------------------- 静态前端 -----------------------------

def _static(name: str):
    return FileResponse(STATIC_DIR / name)


@app.get("/login")
def login_page():
    """独立登录页。"""
    return _static("login.html")


@app.get("/logout")
def logout_page():
    """退出登录页（清 cookie 后跳登录页）。"""
    resp = RedirectResponse(url="/login")
    resp.delete_cookie(auth.COOKIE_NAME, path="/")
    return resp


@app.get("/")
def root_page(request: Request):
    """主界面：已登录返回 index.html，未登录 302 到 /login。"""
    token = request.cookies.get(auth.COOKIE_NAME)
    if token and auth.validate_token(token):
        return _static("index.html")
    return RedirectResponse(url="/login")


@app.get("/style.css")
def style_css():
    return _static("style.css")


@app.get("/app.js")
def app_js():
    return _static("app.js")


@app.get("/login.css")
def login_css():
    return _static("login.css")


@app.get("/login.js")
def login_js():
    return _static("login.js")
